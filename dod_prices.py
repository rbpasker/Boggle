#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""
Deterministic date-of-death (DOD) price lookup for the the estate account workbook.

Rule (Treas. Reg. §20.2031-2(b)): DOD 2025-12-07 is a Sunday; value = mean of the
high–low means on the nearest trading days before and after (2025-12-05, 2025-12-08),
weighted inversely by intervening trading days (zero on each side → simple average).
Mutual funds (§20.2031-8(b)): last NAV quoted on or before the DOD → 2025-12-05 close.

Data: Yahoo Finance chart API, daily bars, one HTTPS request per symbol.  Yahoo's historical
prices are SPLIT-ADJUSTED to today's share count, so split events since the DOD are fetched
in the same request and reported:  F_after_feb = product of split ratios after 2/28/2026
(the statement date whose quantities the CSV carries), F_dod_feb = ratios between DOD and 2/28.
  DOD Price (as traded)      = adj × F_dod_feb × F_after_feb   (per share held on 12/7/2025)
  DOD Price per 2/28 share   = adj × F_after_feb               (→ workbook column V)
  Yahoo import               = adj price, quantity × F_after_feb (today's shares)
Every input is written to the audit CSV.  Run with uv:  uv run dod_prices.py ...

Usage
    python3 dod_prices.py --csv holdings.csv --out dod_prices.csv
        [--yahoo-out A_Yahoo_Portfolio.csv --recipient A]   # Yahoo portfolio import file
        [--xlsx distribution.xlsx]              # fill column V (DOD override)
"""
import argparse, csv, json, ssl, sys, time, urllib.request, urllib.error
from datetime import date, datetime

DOD, BEFORE, AFTER = date(2025, 12, 7), date(2025, 12, 5), date(2025, 12, 8)
FUND_SECTIONS = {"Stock Funds", "Bond Funds", "Short-term Funds"}
YAHOO_MAP = {"BRKB": "BRK-B", "BK": "BNY", "RBGLD": "RBGLY", "756CNT929": None}   # statement ticker -> Yahoo ticker
FIXED = {"FDRXX": 1.0}                                # money market: $1.00 NAV
FEB28 = date(2026, 2, 28)
P1 = 1764720000                                      # 2025-12-03 UTC; window runs to today so split events are included


def yahoo_symbol(sym):
    return YAHOO_MAP.get(sym, sym.replace(".", "-"))


def fetch(sym, ctx=None, retries=4):
    p2 = int(time.time()) + 86400
    url = f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?period1={P1}&period2={p2}&interval=1d&events=splits"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    last = None
    for i in range(retries):
        try:
            with urllib.request.urlopen(req, timeout=20, context=ctx) as r:
                j = json.load(r)
            res = j["chart"]["result"]
            if not res:
                return None
            r0 = res[0]; q = r0["indicators"]["quote"][0]; off = r0["meta"].get("gmtoffset", 0)
            bars = {}
            for k, t in enumerate(r0.get("timestamp") or []):
                if q["close"][k] is None or t > P1 + 8 * 86400:
                    continue
                d = datetime.utcfromtimestamp(t + off).date()
                bars[d] = (q["high"][k], q["low"][k], q["close"][k])
            splits = []
            for v in (r0.get("events", {}).get("splits", {}) or {}).values():
                num, den = v["splitRatio"].split(":")
                splits.append((datetime.utcfromtimestamp(v["date"] + off).date(), float(num) / float(den), v["splitRatio"]))
            return {"bars": bars, "splits": sorted(splits), "type": r0["meta"].get("instrumentType"),
                    "exch": r0["meta"].get("exchangeName"), "last": r0["meta"].get("regularMarketPrice")}
        except urllib.error.HTTPError as e:
            if e.code == 404:
                return None
            last = f"HTTP {e.code}"
        except Exception as e:
            last = repr(e)
        time.sleep(2 * (i + 1))
    raise RuntimeError(last)


def dod_price(bars, is_fund):
    if is_fund:
        prior = [d for d in bars if d <= DOD]
        if not prior:
            return None, "no NAV on or before DOD", {}
        d = max(prior)
        return round(bars[d][2], 4), f"NAV close {d.isoformat()} (§20.2031-8(b))", {"nav_date": d.isoformat(), "nav": bars[d][2]}
    if BEFORE in bars and AFTER in bars:
        h1, l1, _ = bars[BEFORE]; h2, l2, _ = bars[AFTER]
        m1, m2 = (h1 + l1) / 2, (h2 + l2) / 2
        return round((m1 + m2) / 2, 4), "mean of H/L means 12/5 & 12/8 (§20.2031-2(b))", \
            {"high_1205": h1, "low_1205": l1, "mean_1205": m1, "high_1208": h2, "low_1208": l2, "mean_1208": m2}
    for d in (BEFORE, AFTER):
        if d in bars:
            h, l, _ = bars[d]
            return round((h + l) / 2, 4), f"H/L mean {d.isoformat()} only (other side missing)", {f"high_{d:%m%d}": h, f"low_{d:%m%d}": l}
    return None, "no bars on 12/5 or 12/8", {}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", required=True)
    ap.add_argument("--out", default="dod_prices.csv")
    ap.add_argument("--xlsx")
    ap.add_argument("--yahoo-out")
    ap.add_argument("--recipient", default="A")
    ap.add_argument("--sleep", type=float, default=0.3)
    ap.add_argument("--cafile", help="CA bundle for HTTPS (proxy environments)")
    args = ap.parse_args()
    ctx = ssl.create_default_context(cafile=args.cafile) if args.cafile else None

    rows = list(csv.DictReader(open(args.csv, encoding="utf-8-sig")))
    results, seen = [], set()
    for r in rows:
        sym = r["Symbol"]
        if sym in seen:
            continue
        seen.add(sym)
        stmt = float(r["Price 2/28/26"]) if r.get("Price 2/28/26") else None
        ysym = yahoo_symbol(sym)
        rec = {"Symbol": sym, "Yahoo": ysym or "", "Section": r["Section"], "Stmt Price 2/28/26": stmt,
               "DOD Price": None, "DOD Price as traded": None, "DOD Price per 2/28 share": None,
               "F_dod_feb": 1.0, "F_after_feb": 1.0, "Splits since DOD": "", "Method": "", "Flag": "",
               "Yahoo Type": "", "Yahoo Exchange": "", "Yahoo Last": None}
        if ysym is None:
            rec["Method"] = "unpriced (escrow)"
        elif sym in FIXED:
            rec["DOD Price"] = rec["DOD Price as traded"] = rec["DOD Price per 2/28 share"] = FIXED[sym]
            rec["Method"] = "fixed $1.00 NAV (money market)"
        else:
            try:
                data = fetch(ysym, ctx)
                if data is None:
                    rec["Method"] = "symbol not found on Yahoo"; rec["Flag"] = "lookup failed"
                else:
                    rec["Yahoo Type"], rec["Yahoo Exchange"] = data["type"], data["exch"]
                    is_fund = r["Section"] in FUND_SECTIONS or data["type"] == "MUTUALFUND"
                    px, method, detail = dod_price(data["bars"], is_fund)
                    rec["DOD Price"], rec["Method"] = px, method
                    rec["Yahoo Last"] = data["last"]
                    rec.update(detail)
                    f1 = f2 = 1.0
                    for d, ratio, txt in data["splits"]:
                        if DOD < d <= FEB28: f1 *= ratio
                        elif d > FEB28: f2 *= ratio
                    rec["F_dod_feb"], rec["F_after_feb"] = f1, f2
                    rec["Splits since DOD"] = "; ".join(f"{d.isoformat()} {txt}" for d, _, txt in data["splits"] if d > DOD)
                    if px is not None:
                        rec["DOD Price per 2/28 share"] = round(px * f2, 4)
                        rec["DOD Price as traded"] = round(px * f1 * f2, 4)
                        if stmt and not (0.25 * stmt <= px * f2 <= 4 * stmt):
                            rec["Flag"] = f"REVIEW: {round(px * f2, 2)} vs 2/28 {stmt} after split adj — check ticker"
                    else:
                        rec["Flag"] = "no DOD bars"
            except Exception as e:
                rec["Method"] = f"ERROR {e}"; rec["Flag"] = "lookup failed"
            time.sleep(args.sleep)
        results.append(rec)
        print(f"{sym:10s} {str(rec['DOD Price']):>10s}  {rec['Method']}  {rec['Flag']}", file=sys.stderr)

    fields = ["Symbol", "Yahoo", "Section", "Yahoo Type", "Yahoo Exchange", "Stmt Price 2/28/26", "Yahoo Last",
              "DOD Price", "DOD Price per 2/28 share", "DOD Price as traded", "F_dod_feb", "F_after_feb", "Splits since DOD",
              "Method", "Flag", "high_1205", "low_1205", "mean_1205", "high_1208", "low_1208", "mean_1208", "nav_date", "nav"]
    with open(args.out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore"); w.writeheader(); w.writerows(results)
    ok = sum(1 for x in results if x["DOD Price"] is not None and not x["Flag"])
    print(f"\n{ok}/{len(results)} symbols priced without flags -> {args.out}", file=sys.stderr)
    by_sym = {x["Symbol"]: x for x in results}

    if args.yahoo_out:
        qcol = f"{args.recipient} Shares"; n = miss = 0
        with open(args.yahoo_out, "w", newline="") as f:
            w = csv.writer(f); w.writerow(["Symbol", "Trade Date", "Purchase Price", "Quantity"])
            for r in rows:
                q = float(r[qcol] or 0); x = by_sym.get(r["Symbol"])
                if q <= 0 or not x or not x["Yahoo"]:
                    continue
                px = x["DOD Price"] if x["DOD Price"] is not None else ""
                miss += px == ""
                qn = q * x["F_after_feb"]
                w.writerow([x["Yahoo"], DOD.strftime("%Y%m%d"), px, f"{qn:.4f}".rstrip("0").rstrip(".")]); n += 1
        print(f"wrote {n} {args.recipient} positions to {args.yahoo_out} ({miss} without a DOD price)", file=sys.stderr)

    if args.xlsx:
        from openpyxl import load_workbook
        wb = load_workbook(args.xlsx); ws = wb["Holdings"]
        hdr = {ws.cell(row=5, column=c).value: c for c in range(1, ws.max_column + 1)}
        c_sym, c_ovr = hdr["Symbol"], hdr["DOD Price Override (manual)"]; n = 0
        for rr in range(6, ws.max_row + 1):
            x = by_sym.get(ws.cell(row=rr, column=c_sym).value)
            if x and x["DOD Price per 2/28 share"] is not None:
                ws.cell(row=rr, column=c_ovr, value=x["DOD Price per 2/28 share"]); n += 1
        wb.calculation.fullCalcOnLoad = True; wb.save(args.xlsx)
        print(f"wrote {n} DOD prices into column V of {args.xlsx}", file=sys.stderr)


if __name__ == "__main__":
    main()
